"""Evaluates every formula in docs/review/model.xlsx with the 'formulas' engine, reports errors, prints key outputs."""
import json
import sys

import formulas
from openpyxl import load_workbook

path = 'docs/review/model.xlsx'
xl = formulas.ExcelModel().loads(path).finish()
sol = xl.calculate()

wb = load_workbook(path)
values, errors, n_formulas = {}, [], 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                n_formulas += 1
                key = f"'[model.xlsx]{ws.title.upper()}'!{c.coordinate}"
                v = sol.get(key)
                val = v.value[0][0] if v is not None and hasattr(v, 'value') else v
                values[f'{ws.title}!{c.coordinate}'] = val
                s = str(val)
                if v is None or s.startswith('#') or 'Error' in type(val).__name__ or s in ('#NAME?', '#VALUE!', '#DIV/0!', '#N/A', '#REF!'):
                    errors.append((f'{ws.title}!{c.coordinate}', s, c.value[:80]))
print(json.dumps({'total_formulas': n_formulas, 'total_errors': len(errors), 'errors': errors[:20]}, indent=1, default=str))
sc = wb['Scenarios']
for r in range(4, sc.max_row + 1):
    label = sc.cell(row=r, column=1).value
    if label:
        vals = [values.get(f'Scenarios!{col}{r}') for col in 'BCD']
        print(f'{label:52} ' + ' | '.join(f'{v:,.2f}' if isinstance(v, (int, float)) else str(v) for v in vals))
json.dump({k: (v if isinstance(v, (int, float, str)) else str(v)) for k, v in values.items()}, open('docs/review/tools/synthesis/model-values.json', 'w'))
sys.exit(1 if errors else 0)
